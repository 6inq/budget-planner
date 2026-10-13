"""Tests for per-bill, per-month "Paid" marks on the Expenses sheet.

Requirements:
  R1  Every month has a Paid column with a dropdown offering "Paid", for every expense input row.
  R2  `paid` in the config puts "Paid" in exactly the matching month cells; omitted means none.
  R3  Paid marks never change any total (summary JSON, TOTAL row, Dashboard ranges).
  R4  PAID and LEFT TO PAY rows under TOTAL sum the marked and unmarked amounts per month and per year.
  R5  A marked month turns its amount cell and Paid cell green.
  R6  read_workbook.py reads the marks back, so updates keep them.
  R7  Invalid `paid` values are rejected with a clear error.
  R8  Business workbooks get the same columns; the payroll formula row has no input marks.

Run from the repository root:
    python -m unittest discover -s tests -v
"""

import importlib.util
import json
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

SCRIPTS = Path(__file__).resolve().parents[1] / "skills" / "budget-planner" / "scripts"
sys.path.insert(0, str(SCRIPTS))


def load(name: str):
    spec = importlib.util.spec_from_file_location(name, SCRIPTS / f"{name}.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


bb = load("build_budget")
rw = load("read_workbook")

FIRST_ONLY = [True] + [False] * 11


def config(expenses, **overrides) -> dict:
    raw = {"year": 2026, "start_month": 4, "pay": {"frequency": "monthly", "amount": 4000}, "expenses": expenses}
    raw.update(overrides)
    return raw


def rent(paid=None) -> dict:
    item = {"name": "Rent", "category": "Housing", "type": "Bill", "amount": 1200, "frequency": "monthly", "due_day": 1}
    if paid is not None:
        item["paid"] = paid
    return item


def gas() -> dict:
    return {"name": "Gas", "category": "Transportation", "type": "Expense", "months": [80] + [0] * 11}


class PaidCase(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def build(self, raw: dict, name="out.xlsx"):
        result = bb.build_from_config(raw, str(self.tmp / name))
        return result, load_workbook(result["file"])

    @staticmethod
    def paid_cell(ws, row: int, month: int):
        return ws.cell(row=row, column=bb.PAID_COL + month)

    @staticmethod
    def total_row(ws) -> int:
        col = get_column_letter(bb.EXPENSE_MONTH_COL)
        for row in range(bb.FIRST_ROW, ws.max_row + 1):
            if str(ws.cell(row=row, column=bb.EXPENSE_MONTH_COL).value).startswith(f"=SUM({col}{bb.FIRST_ROW}:"):
                return row
        raise AssertionError("no TOTAL row")


class TestColumnsAndDropdowns(PaidCase):  # R1
    def test_twelve_paid_headers_follow_the_budget_window(self):
        _, wb = self.build(config([rent()]))
        ws = wb["Expenses"]
        headers = [self.paid_cell(ws, bb.HEADER_ROW, i) for i in range(12)]
        self.assertEqual(headers[0].value.strftime("%Y-%m"), "2026-04")
        self.assertEqual(headers[11].value.strftime("%Y-%m"), "2027-03")
        self.assertTrue(all(h.number_format == '"Paid" mmm' for h in headers))

    def test_dropdown_offers_paid_on_every_input_row(self):
        _, wb = self.build(config([rent()]))
        ws = wb["Expenses"]
        last = self.total_row(ws) - 1
        expected = f"{get_column_letter(bb.PAID_COL)}{bb.FIRST_ROW}:{get_column_letter(bb.PAID_COL + 11)}{last}"
        matches = [dv for dv in ws.data_validations.dataValidation if dv.formula1 == '"Paid"']
        self.assertEqual(len(matches), 1)
        self.assertEqual(str(matches[0].sqref), expected)

    def test_paid_columns_do_not_overlap_amounts_or_annual_total(self):
        self.assertGreater(bb.PAID_COL, bb.EXPENSE_TOTAL_COL + 1)
        self.assertGreaterEqual(bb.SHEET_FILL_COLS, bb.PAID_COL + 11)


class TestMarks(PaidCase):  # R2
    def test_marks_land_only_in_matching_months(self):
        _, wb = self.build(config([rent([True, False, True] + [False] * 9), gas()]))
        ws = wb["Expenses"]
        rent_row, gas_row = bb.FIRST_ROW, bb.FIRST_ROW + 1
        self.assertEqual([self.paid_cell(ws, rent_row, i).value for i in range(3)], ["Paid", None, "Paid"])
        self.assertTrue(all(self.paid_cell(ws, rent_row, i).value is None for i in range(3, 12)))
        self.assertTrue(all(self.paid_cell(ws, gas_row, i).value is None for i in range(12)))

    def test_all_months_paid(self):
        _, wb = self.build(config([rent([True] * 12)]))
        ws = wb["Expenses"]
        self.assertTrue(all(self.paid_cell(ws, bb.FIRST_ROW, i).value == "Paid" for i in range(12)))

    def test_empty_budget_has_no_marks(self):
        _, wb = self.build(config([]))
        ws = wb["Expenses"]
        last = self.total_row(ws) - 1
        marks = [self.paid_cell(ws, r, i).value for r in range(bb.FIRST_ROW, last + 1) for i in range(12)]
        self.assertEqual(set(marks), {None})


class TestTotalsUnchanged(PaidCase):  # R3
    def test_summary_identical_with_and_without_marks(self):
        unpaid, _ = self.build(config([rent(), gas()]), "a.xlsx")
        paid, _ = self.build(config([rent([True] * 12), gas()]), "b.xlsx")
        for key in ("annual_income", "annual_out", "net", "monthly"):
            self.assertEqual(unpaid[key], paid[key])

    def test_total_row_still_sums_amounts_only(self):
        _, wb = self.build(config([rent(FIRST_ONLY)]))
        ws = wb["Expenses"]
        total = self.total_row(ws)
        self.assertEqual(ws.cell(row=total, column=bb.EXPENSE_MONTH_COL).value, f"=SUM(F{bb.FIRST_ROW}:F{total - 1})")


class TestPaidTotals(PaidCase):  # R4
    def test_paid_and_left_rows_under_total(self):
        _, wb = self.build(config([rent(FIRST_ONLY)]))
        ws = wb["Expenses"]
        total, last = self.total_row(ws), self.total_row(ws) - 1
        self.assertEqual(ws.cell(row=total + 1, column=2).value, "PAID")
        self.assertEqual(ws.cell(row=total + 2, column=2).value, "LEFT TO PAY")
        self.assertEqual(ws.cell(row=total + 1, column=bb.EXPENSE_MONTH_COL).value,
                         f'=SUMIFS(F{bb.FIRST_ROW}:F{last},T{bb.FIRST_ROW}:T{last},"Paid")')
        self.assertEqual(ws.cell(row=total + 2, column=bb.EXPENSE_MONTH_COL).value, f"=F{total}-F{total + 1}")
        sep = get_column_letter(bb.EXPENSE_MONTH_COL + 11)
        self.assertEqual(ws.cell(row=total + 1, column=bb.EXPENSE_MONTH_COL + 11).value,
                         f'=SUMIFS({sep}{bb.FIRST_ROW}:{sep}{last},AE{bb.FIRST_ROW}:AE{last},"Paid")')
        self.assertEqual(ws.cell(row=total + 2, column=bb.EXPENSE_TOTAL_COL).value, f"=SUM(F{total + 2}:Q{total + 2})")

    def test_print_area_includes_paid_columns_and_rows(self):
        _, wb = self.build(config([rent()]))
        ws = wb["Expenses"]
        self.assertIn(f"$AF${self.total_row(ws) + 3}", str(ws.print_area))


class TestHighlight(PaidCase):  # R5
    def test_amounts_and_marks_turn_green(self):
        _, wb = self.build(config([rent()]))
        ws = wb["Expenses"]
        rules = {str(cf.sqref): cf.rules for cf in ws.conditional_formatting}
        last = self.total_row(ws) - 1
        for cell_range in (f"F{bb.FIRST_ROW}:Q{last}", f"T{bb.FIRST_ROW}:AE{last}"):
            self.assertIn(cell_range, rules)
            rule = rules[cell_range][0]
            self.assertEqual(rule.formula, [f'T{bb.FIRST_ROW}="Paid"'])
            self.assertTrue(rule.dxf.font.color.rgb.endswith(bb.POSITIVE))


class TestRoundTrip(PaidCase):  # R6
    def test_marks_survive_read_and_rebuild(self):
        marks = [True, True] + [False] * 9 + [True]
        result, _ = self.build(config([rent(marks), gas()]))
        read = rw.read_workbook(result["file"])
        by_name = {e["name"]: e for e in read["expenses"]}
        self.assertEqual(by_name["Rent"]["paid"], marks)
        self.assertEqual(by_name["Gas"]["paid"], [False] * 12)
        rebuilt, wb = self.build(read, "rebuilt.xlsx")
        self.assertEqual(self.paid_cell(wb["Expenses"], bb.FIRST_ROW, 11).value, "Paid")
        self.assertEqual(rebuilt["annual_out"], result["annual_out"])

    def test_mark_typed_in_excel_any_case_is_read(self):
        result, wb = self.build(config([rent()]))
        ws = wb["Expenses"]
        self.paid_cell(ws, bb.FIRST_ROW, 0).value = "paid"
        self.paid_cell(ws, bb.FIRST_ROW, 1).value = "  "
        wb.save(result["file"])
        self.assertEqual(rw.read_workbook(result["file"])["expenses"][0]["paid"], FIRST_ONLY)

    def test_workbook_without_paid_columns_reads_as_unpaid(self):
        result, wb = self.build(config([rent()]))
        ws = wb["Expenses"]
        for row in range(bb.HEADER_ROW, ws.max_row + 1):
            for i in range(12):
                self.paid_cell(ws, row, i).value = None
        wb.save(result["file"])
        self.assertEqual(rw.read_workbook(result["file"])["expenses"][0]["paid"], [False] * 12)


class TestValidation(PaidCase):  # R7
    def test_bad_paid_values_rejected(self):
        for bad in ([True] * 11, [True] * 13, ["Paid"] + [False] * 11, [1] + [False] * 11, "Paid", {}):
            with self.subTest(bad=bad), self.assertRaisesRegex(ValueError, "12 true/false"):
                bb.validate_config(config([rent(bad)]))

    def test_cli_reports_error_for_bad_paid(self):
        path = self.tmp / "bad.json"
        path.write_text(json.dumps(config([rent([True])])), encoding="utf-8")
        self.assertEqual(bb.main(["--config", str(path), "--out", str(self.tmp / "x.xlsx")]), 2)


class TestBusiness(PaidCase):  # R8
    def test_business_paid_block_skips_payroll_row(self):
        item = {"name": "Office Rent", "category": "Rent", "type": "Fixed", "amount": 900, "paid": FIRST_ONLY}
        result, wb = self.build(config([item], budget_type="business"))
        ws = wb["Expenses"]
        self.assertIsNone(self.paid_cell(ws, bb.FIRST_ROW, 0).value)
        self.assertEqual(self.paid_cell(ws, bb.FIRST_ROW + 1, 0).value, "Paid")
        read = rw.read_workbook(result["file"])
        self.assertEqual(read["expenses"][0]["paid"], FIRST_ONLY)


if __name__ == "__main__":
    unittest.main()
