"""Tests for the optional Loan Payoff Plan sheet.

Requirements:
  P1  The sheet exists only when features.payoff_plan is on, placed after the trackers and before Lists.
  P2  Every input is written next to its label; paychecks per month defaults from the pay frequency.
  P3  The extra-payment goal is take-home minus bills, loan payment, savings and buffer, never below 0.
  P4  The schedule has PLAN_MONTHS rows starting at the plan start date and the loan balance.
  P5  An Extra Paid amount replaces that month's goal; blank uses the goal.
  P6  Status and Underwater By are colored: underwater red, above water / paid off green.
  P7  read_workbook.py reads inputs and Extra Paid back, so updates keep them.
  P8  Invalid plan values are rejected with a clear error.
  P9  A plan with no numbers still builds and shows blanks, not errors.
  P10 The Dashboard is unchanged by the plan; business workbooks get the same sheet.
  P13 Car or house loans: house gets a 360-month schedule, home wording, and a value that may rise.
  P14 The questionnaire asks about a Payoff Plan and its loan details before the last (fill-out) question.
  P12 The schedule stops at the payoff month; the no-extra dates still cover all months via hidden helpers.
  P11 An optional underwater target shows the month it is reached and highlights that row; blank when unset.

Run from the repository root:
    python -m unittest discover -s tests -v
"""

import importlib.util
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

SCRIPTS = Path(__file__).resolve().parents[1] / "skills" / "budget-planner" / "scripts"
sys.path.insert(0, str(SCRIPTS))


def load(name: str):
    spec = importlib.util.spec_from_file_location(name, SCRIPTS / f"{name}.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


bb = load("build_budget")
rw = load("read_workbook")

PLAN = {"loan_name": "Car", "balance": 18000, "apr": 6.5, "payment": 420, "asset_value": 12000,
        "value_drop": 120, "take_home": 1900, "paychecks_per_month": 2, "monthly_bills": 2100,
        "savings_percent": 15, "buffer": 100, "start": "2027-02-01"}
FIRST = bb.PLAN_HEADER_ROW + 1
LAST = bb.PLAN_HEADER_ROW + bb.PLAN_MONTHS
GOAL_ROW = bb.PLAN_INPUT_ROW + bb.PLAN_GOAL_OFFSET


def config(plan=None, enabled=True, **overrides) -> dict:
    raw = {"year": 2027, "pay": {"frequency": "biweekly", "amount": 1900, "first_payday": "2027-01-08"},
           "features": {"payoff_plan": enabled}}
    if plan is not None:
        raw["payoff_plan"] = plan
    raw.update(overrides)
    return raw


class PlanCase(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def build(self, raw: dict, name="out.xlsx"):
        result = bb.build_from_config(raw, str(self.tmp / name))
        return result, load_workbook(result["file"])

    @staticmethod
    def value(ws, key: str):
        return ws.cell(row=bb.plan_row(key), column=3).value


class TestSheetPresence(PlanCase):  # P1
    def test_sheet_only_when_enabled(self):
        _, off = self.build(config(PLAN, enabled=False), "off.xlsx")
        _, on = self.build(config(PLAN), "on.xlsx")
        self.assertNotIn("Payoff Plan", off.sheetnames)
        self.assertIn("Payoff Plan", on.sheetnames)

    def test_sheet_order_after_trackers_before_lists(self):
        raw = config(PLAN)
        raw["features"].update(savings_goals=True, debt_tracker=True)
        _, wb = self.build(raw)
        names = wb.sheetnames
        self.assertEqual(names.index("Payoff Plan"), names.index("Debt Tracker") + 1)
        self.assertEqual(names.index("Lists"), names.index("Payoff Plan") + 1)


class TestInputs(PlanCase):  # P2
    def test_every_input_written_beside_its_label(self):
        _, wb = self.build(config(PLAN))
        ws = wb["Payoff Plan"]
        for key, label, _ in bb.PLAN_INPUTS:
            self.assertEqual(ws.cell(row=bb.plan_row(key), column=2).value, label)
        self.assertEqual(self.value(ws, "loan_name"), "Car")
        self.assertEqual(self.value(ws, "balance"), 18000)
        self.assertEqual(self.value(ws, "apr"), 6.5)
        self.assertEqual(self.value(ws, "start").date(), date(2027, 2, 1))

    def test_paychecks_per_month_defaults_from_frequency(self):
        expected = {"weekly": 52 / 12, "biweekly": 26 / 12, "semimonthly": 2, "monthly": 1}
        for frequency, per_month in expected.items():
            pay = {"frequency": frequency, "amount": 0}
            if frequency == "semimonthly":
                pay["semimonthly_days"] = ["1", "15"]
            plan = {k: v for k, v in PLAN.items() if k != "paychecks_per_month"}
            with self.subTest(frequency=frequency):
                _, wb = self.build(config(plan, pay=pay), f"{frequency}.xlsx")
                self.assertAlmostEqual(self.value(wb["Payoff Plan"], "paychecks_per_month"), per_month)

    def test_explicit_paychecks_per_month_wins(self):
        _, wb = self.build(config({**PLAN, "paychecks_per_month": 3}))
        self.assertEqual(self.value(wb["Payoff Plan"], "paychecks_per_month"), 3)


class TestGoal(PlanCase):  # P3
    def test_goal_formula_subtracts_everything_and_floors_at_zero(self):
        _, wb = self.build(config(PLAN))
        ws = wb["Payoff Plan"]
        r = bb.PLAN_INPUT_ROW
        self.assertEqual(ws.cell(row=GOAL_ROW, column=5).value, "EXTRA PAYMENT GOAL (per month)")
        self.assertEqual(ws.cell(row=GOAL_ROW, column=7).value, f"=MAX(0,ROUND(G{r}-G{r + 1}-G{r + 2}-G{r + 3},2))")
        c = {k: f"$C${bb.plan_row(k)}" for k, _, _ in bb.PLAN_INPUTS}
        self.assertEqual(ws[f"G{r}"].value, f"=N({c['take_home']})*N({c['paychecks_per_month']})")
        self.assertEqual(ws[f"G{r + 1}"].value, f"=N({c['monthly_bills']})+N({c['payment']})")
        self.assertEqual(ws[f"G{r + 2}"].value, f"=ROUND(G{r}*N({c['savings_percent']})/100,2)")
        self.assertEqual(ws[f"G{r + 3}"].value, f"=N({c['buffer']})")


class TestSchedule(PlanCase):  # P4, P5
    def test_schedule_length_and_first_row(self):
        _, wb = self.build(config(PLAN))
        ws = wb["Payoff Plan"]
        c = {k: f"$C${bb.plan_row(k)}" for k, _, _ in bb.PLAN_INPUTS}
        self.assertEqual(ws.cell(row=FIRST, column=2).value, f'=IF({c["start"]}="","",EDATE({c["start"]},0))')
        self.assertEqual(ws.cell(row=LAST, column=2).value,
                         f'=IF(N(H{LAST - 1})>0,IF({c["start"]}="","",EDATE({c["start"]},{bb.PLAN_MONTHS - 1})),"")')
        self.assertEqual(ws.cell(row=FIRST, column=3).value, f"=N({c['balance']})")
        self.assertEqual(ws.cell(row=FIRST + 1, column=3).value, f'=IF(N(H{FIRST})>0,H{FIRST},"")')
        self.assertIsNone(ws.cell(row=LAST + 1, column=2).value)

    def test_extra_paid_overrides_goal_and_is_written(self):
        _, wb = self.build(config({**PLAN, "extra_paid": [300, None, 0]}))
        ws = wb["Payoff Plan"]
        self.assertEqual([ws.cell(row=FIRST + n, column=7).value for n in range(4)], [300, None, 0, None])
        self.assertEqual(ws.cell(row=FIRST, column=8).value,
                         f'=MAX(0,C{FIRST}+D{FIRST}-E{FIRST}-IF(G{FIRST}="",F{FIRST},G{FIRST}))')
        self.assertEqual(ws.cell(row=FIRST, column=6).value,
                         f"=MIN($G${GOAL_ROW},MAX(0,C{FIRST}+D{FIRST}-E{FIRST}))")


class TestColors(PlanCase):  # P6
    def test_status_and_gap_rules(self):
        _, wb = self.build(config(PLAN))
        ws = wb["Payoff Plan"]
        rules = {}
        for cf in ws.conditional_formatting:
            rules.setdefault(str(cf.sqref), []).extend(cf.rules)
        status = {r.formula[0]: r.dxf.font.color.rgb[-6:] for r in rules[f"K{FIRST}:K{LAST}"]}
        self.assertEqual(status[f'K{FIRST}="Underwater"'], bb.NEGATIVE)
        self.assertEqual(status[f'K{FIRST}="Above water"'], bb.POSITIVE)
        self.assertEqual(status[f'K{FIRST}="Paid off"'], bb.POSITIVE)
        gap = {r.operator: r.dxf.font.color.rgb[-6:] for r in rules[f"J{FIRST}:J{LAST}"]}
        self.assertEqual(gap, {"greaterThan": bb.NEGATIVE, "lessThanOrEqual": bb.POSITIVE})


class TestRoundTrip(PlanCase):  # P7
    def test_inputs_and_extra_paid_survive(self):
        result, _ = self.build(config({**PLAN, "extra_paid": [250, None, 400]}))
        read = rw.read_workbook(result["file"])
        self.assertTrue(read["features"]["payoff_plan"])
        plan = read["payoff_plan"]
        for key in ("balance", "apr", "payment", "asset_value", "value_drop", "take_home", "monthly_bills",
                    "savings_percent", "buffer", "paychecks_per_month"):
            self.assertEqual(plan[key], PLAN[key], key)
        self.assertEqual(plan["loan_name"], "Car")
        self.assertEqual(plan["start"], "2027-02-01")
        self.assertEqual(plan["extra_paid"], [250, None, 400])
        _, wb = self.build(read, "rebuilt.xlsx")
        self.assertEqual(wb["Payoff Plan"].cell(row=FIRST + 2, column=7).value, 400)

    def test_amounts_typed_in_excel_are_read(self):
        result, wb = self.build(config(PLAN))
        ws = wb["Payoff Plan"]
        ws.cell(row=bb.plan_row("apr"), column=3).value = 8.9
        ws.cell(row=FIRST + 5, column=7).value = 650
        wb.save(result["file"])
        plan = rw.read_workbook(result["file"])["payoff_plan"]
        self.assertEqual(plan["apr"], 8.9)
        self.assertEqual(plan["extra_paid"], [None] * 5 + [650])

    def test_workbook_without_plan_reads_as_off(self):
        result, _ = self.build(config(PLAN, enabled=False))
        read = rw.read_workbook(result["file"])
        self.assertFalse(read["features"]["payoff_plan"])
        self.assertNotIn("payoff_plan", read)


class TestValidation(PlanCase):  # P8
    def test_bad_values_rejected(self):
        bad = [({"balance": -1}, "cannot be negative"), ({"apr": "ten"}, "must be a number"),
               ({"start": "next month"}, "must be a date"),
               ({"extra_paid": [1] * (bb.PLAN_MONTHS + 1)}, "at most"),
               ({"extra_paid": "lots"}, "at most"), ({"extra_paid": [-5]}, "cannot be negative")]
        for change, message in bad:
            with self.subTest(change=change), self.assertRaisesRegex(ValueError, message):
                bb.validate_config(config({**PLAN, **change}))


class TestBlankPlan(PlanCase):  # P9
    def test_enabled_without_numbers_builds_with_blank_results(self):
        result, wb = self.build(config())
        ws = wb["Payoff Plan"]
        self.assertIsNone(self.value(ws, "balance"))
        self.assertAlmostEqual(self.value(ws, "paychecks_per_month"), 26 / 12)
        self.assertTrue(ws.cell(row=FIRST, column=11).value.startswith('=IF(N($C$'))
        self.assertTrue(ws.cell(row=bb.PLAN_INPUT_ROW + 5, column=7).value.startswith('=IF(N($C$'))
        read = rw.read_workbook(result["file"])
        self.assertEqual(read["payoff_plan"]["extra_paid"], [])


class TestOtherSheets(PlanCase):  # P10
    def test_dashboard_unchanged_by_plan(self):
        _, off = self.build(config(PLAN, enabled=False), "off.xlsx")
        _, on = self.build(config(PLAN), "on.xlsx")
        self.assertEqual(str(off["Dashboard"].print_area), str(on["Dashboard"].print_area))
        for row in range(1, 60):
            for col in range(1, 16):
                self.assertEqual(off["Dashboard"].cell(row=row, column=col).value,
                                 on["Dashboard"].cell(row=row, column=col).value)

    def test_business_workbook_gets_plan(self):
        result, wb = self.build(config(PLAN, budget_type="business"))
        self.assertIn("Payoff Plan", wb.sheetnames)
        self.assertEqual(rw.read_workbook(result["file"])["payoff_plan"]["balance"], 18000)


class TestTarget(PlanCase):  # P11
    TARGET_ROW = bb.PLAN_INPUT_ROW + 9

    def test_target_month_and_highlight(self):
        _, wb = self.build(config({**PLAN, "target_gap": 3000}))
        ws = wb["Payoff Plan"]
        target = f"$C${bb.plan_row('target_gap')}"
        self.assertEqual(self.value(ws, "target_gap"), 3000)
        self.assertEqual(ws.cell(row=self.TARGET_ROW, column=5).value, "Target reached by")
        formula = ws.cell(row=self.TARGET_ROW, column=7).value
        self.assertTrue(formula.startswith(f"=IF(ISNUMBER({target}),IF(N($C$"))
        self.assertIn(f"INDEX($J${FIRST}:$J${LAST}<={target},0)", formula)
        rules = [r for cf in ws.conditional_formatting if str(cf.sqref) == f"B{FIRST}:L{LAST}" for r in cf.rules
                 if "ISNUMBER" in r.formula[0]]
        self.assertEqual(rules[0].formula, [f"AND(ISNUMBER({target}),$J{FIRST}<={target},"
                                            f"OR(ROW()={FIRST},N($J{FIRST - 1})>{target}))"])

    def test_no_target_leaves_input_blank(self):
        _, wb = self.build(config(PLAN))
        self.assertIsNone(self.value(wb["Payoff Plan"], "target_gap"))

    def test_zero_target_means_above_water(self):
        result, _ = self.build(config({**PLAN, "target_gap": 0}))
        self.assertEqual(rw.read_workbook(result["file"])["payoff_plan"]["target_gap"], 0)

    def test_target_round_trip_and_validation(self):
        result, _ = self.build(config({**PLAN, "target_gap": 3500}))
        self.assertEqual(rw.read_workbook(result["file"])["payoff_plan"]["target_gap"], 3500)
        with self.assertRaisesRegex(ValueError, "cannot be negative"):
            bb.validate_config(config({**PLAN, "target_gap": -1}))


class TestStopsAtPayoff(PlanCase):  # P12
    def test_rows_after_payoff_are_blank_formulas(self):
        _, wb = self.build(config(PLAN))
        ws = wb["Payoff Plan"]
        for col in range(2, 13):
            if col == 7:
                continue
            with self.subTest(col=col):
                formula = ws.cell(row=FIRST + 1, column=col).value
                self.assertTrue(formula.startswith(f"=IF(N(H{FIRST})>0,"), formula)
                self.assertTrue(formula.endswith(',"")'), formula)
        self.assertFalse(ws.cell(row=FIRST, column=3).value.startswith("=IF(N(H"))

    def test_helper_columns_hidden_and_full_length(self):
        _, wb = self.build(config(PLAN))
        ws = wb["Payoff Plan"]
        for i in range(3):
            letter = get_column_letter(bb.PLAN_HELPER_COL + i)
            self.assertTrue(ws.column_dimensions[letter].hidden, letter)
            self.assertFalse(str(ws.cell(row=LAST, column=bb.PLAN_HELPER_COL + i).value).startswith("=IF(N(H"))
        self.assertEqual(ws.cell(row=FIRST, column=12).value, f"=O{FIRST}")

    def test_no_extra_results_use_helpers(self):
        _, wb = self.build(config(PLAN))
        ws = wb["Payoff Plan"]
        above, paid = (ws.cell(row=bb.PLAN_INPUT_ROW + i, column=7).value for i in (6, 8))
        self.assertIn(f"INDEX($N${FIRST}:$N${LAST},MATCH(TRUE,INDEX($O${FIRST}:$O${LAST}-$P${FIRST}:$P${LAST}<=0,0)", above)
        self.assertIn(f"INDEX($N${FIRST}:$N${LAST},MATCH(TRUE,INDEX($O${FIRST}:$O${LAST}<=0,0)", paid)

    def test_blank_rows_lose_stripes(self):
        _, wb = self.build(config(PLAN))
        ws = wb["Payoff Plan"]
        rules = [r for cf in ws.conditional_formatting if str(cf.sqref) == f"B{FIRST}:L{LAST}" for r in cf.rules]
        blank = [r for r in rules if r.formula == [f'$C{FIRST}=""']]
        self.assertEqual(len(blank), 1)
        self.assertTrue(blank[0].stopIfTrue)


class TestLoanTypes(PlanCase):  # P13
    HOUSE = {**PLAN, "loan_type": "house", "balance": 240000, "apr": 6.25, "payment": 1500,
             "asset_value": 260000, "value_drop": -400}

    def test_car_is_default(self):
        _, wb = self.build(config(PLAN))
        ws = wb["Payoff Plan"]
        self.assertEqual(ws["B2"].value, "Car Loan Payoff Plan")
        self.assertEqual(ws.cell(row=bb.PLAN_HEADER_ROW, column=9).value, "Vehicle Value")
        self.assertEqual(ws.cell(row=bb.plan_row("asset_value"), column=2).value, "Vehicle value today")
        self.assertIsNotNone(ws.cell(row=LAST, column=2).value)
        self.assertIsNone(ws.cell(row=LAST + 1, column=2).value)

    def test_house_wording_and_360_month_schedule(self):
        _, wb = self.build(config(self.HOUSE))
        ws = wb["Payoff Plan"]
        last = bb.PLAN_HEADER_ROW + 360
        self.assertEqual(ws["B2"].value, "Home Loan Payoff Plan")
        self.assertEqual(ws.cell(row=bb.PLAN_HEADER_ROW, column=9).value, "Home Value")
        self.assertEqual(ws.cell(row=bb.plan_row("payment"), column=2).value, "Regular payment (principal + interest)")
        self.assertEqual(ws.cell(row=bb.plan_row("value_drop"), column=2).value,
                         "Value change per month (negative if rising)")
        self.assertIsNotNone(ws.cell(row=last, column=2).value)
        self.assertIsNone(ws.cell(row=last + 1, column=2).value)
        self.assertIn('"After 30 years"', ws.cell(row=bb.PLAN_INPUT_ROW + 5, column=7).value)
        self.assertIn(f"$J${FIRST}:$J${last}", ws.cell(row=bb.PLAN_INPUT_ROW + 5, column=7).value)

    def test_rising_value_allowed_only_for_value_change(self):
        self.assertEqual(bb.validate_config(config(self.HOUSE))["payoff_plan"]["value_drop"], -400)
        with self.assertRaisesRegex(ValueError, "cannot be negative"):
            bb.validate_config(config({**self.HOUSE, "asset_value": -1}))

    def test_house_round_trip_keeps_type_and_long_extra_list(self):
        extra = [None] * 100 + [250]
        result, _ = self.build(config({**self.HOUSE, "extra_paid": extra}))
        plan = rw.read_workbook(result["file"])["payoff_plan"]
        self.assertEqual(plan["loan_type"], "house")
        self.assertEqual(plan["value_drop"], -400)
        self.assertEqual(plan["extra_paid"], extra)

    def test_extra_list_limit_follows_type(self):
        bb.validate_config(config({**self.HOUSE, "extra_paid": [1] * 360}))
        with self.assertRaisesRegex(ValueError, "at most 72"):
            bb.validate_config(config({**PLAN, "extra_paid": [1] * 73}))
        with self.assertRaisesRegex(ValueError, "at most 360"):
            bb.validate_config(config({**self.HOUSE, "extra_paid": [1] * 361}))

    def test_bad_loan_type_rejected(self):
        with self.assertRaisesRegex(ValueError, "loan_type"):
            bb.validate_config(config({**PLAN, "loan_type": "boat"}))


class TestQuestionnaire(unittest.TestCase):  # P14
    def test_payoff_question_before_last_question(self):
        text = (SCRIPTS.parent / "SKILL.md").read_text(encoding="utf-8")
        round3 = text[text.index("**Round 3**"):text.index("## Step 1b")]
        ask = round3.index("Do you want a Payoff Plan sheet to pay off a car or home loan faster?")
        details = round3.index("Loan balance, interest rate (APR), and regular monthly payment")
        last = round3.index('LAST question: "Do you want to fill out the sheet, or me?"')
        self.assertLess(ask, details)
        self.assertLess(details, last)
        self.assertEqual(round3.count("LAST question"), 1)


if __name__ == "__main__":
    unittest.main()
