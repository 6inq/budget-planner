"""Requirement-driven tests for the Budget Planner generator.

Run from the repository root:
    python -m unittest discover -s tests -v
"""

import copy
import importlib.util
import json
import subprocess
import sys
import tempfile
import unittest
from datetime import date, timedelta
from pathlib import Path

from openpyxl import load_workbook

SCRIPTS = Path(__file__).resolve().parents[1] / "skills" / "budget-planner" / "scripts"
SCRIPT = SCRIPTS / "build_budget.py"
spec = importlib.util.spec_from_file_location("build_budget", SCRIPT)
bb = importlib.util.module_from_spec(spec)
spec.loader.exec_module(bb)


def base(**overrides) -> dict:
    raw = {"year": 2026, "pay": {"frequency": "monthly", "amount": 0}}
    raw.update(overrides)
    return raw


def paydays(**pay) -> list[date]:
    return bb.generate_paydays(bb.validate_config(base(pay=pay)))


class TempDirCase(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def build(self, raw: dict, name="out.xlsx"):
        result = bb.build_from_config(raw, str(self.tmp / name))
        return result, load_workbook(result["file"])


# ---------------------------------------------------------------- pay schedules

class PayScheduleTests(unittest.TestCase):
    def test_weekly_count_and_spacing(self):
        days = paydays(frequency="weekly", amount=500)
        self.assertIn(len(days), (52, 53))
        self.assertTrue(all((b - a).days == 7 for a, b in zip(days, days[1:])))

    def test_biweekly_count_and_spacing(self):
        days = paydays(frequency="biweekly", amount=1000, first_payday="2026-01-09")
        self.assertIn(len(days), (26, 27))
        self.assertEqual(days[0], date(2026, 1, 9))
        self.assertTrue(all((b - a).days == 14 for a, b in zip(days, days[1:])))

    def test_first_payday_before_window_rolls_forward(self):
        days = paydays(frequency="biweekly", first_payday="2025-12-26")
        self.assertEqual(days[0], date(2026, 1, 9))

    def test_first_payday_late_in_window_back_fills_cadence(self):
        days = paydays(frequency="weekly", first_payday="2026-03-06")
        self.assertEqual(days[0], date(2026, 1, 2))

    def test_default_first_payday_is_first_friday(self):
        days = paydays(frequency="weekly")
        self.assertEqual(days[0], date(2026, 1, 2))
        self.assertEqual(days[0].weekday(), 4)

    def test_semimonthly_first_and_fifteenth(self):
        days = paydays(frequency="semimonthly", semimonthly_days=["15", "1"])
        self.assertEqual(len(days), 24)
        self.assertEqual({d.day for d in days}, {1, 15})

    def test_semimonthly_fifteenth_and_last_day_non_leap(self):
        days = paydays(frequency="semimonthly", semimonthly_days=["15", "last"])
        self.assertIn(date(2026, 2, 28), days)
        self.assertIn(date(2026, 4, 30), days)
        self.assertIn(date(2026, 12, 31), days)

    def test_semimonthly_last_day_leap_year(self):
        config = bb.validate_config({"year": 2028, "pay": {"frequency": "semimonthly",
                                                           "semimonthly_days": ["1", "last"]}})
        days = bb.generate_paydays(config)
        self.assertIn(date(2028, 2, 29), days)
        self.assertIn(date(2028, 2, 1), days)

    def test_monthly_last_and_fifteenth(self):
        self.assertEqual(paydays(frequency="monthly", monthly_day="last")[1], date(2026, 2, 28))
        self.assertTrue(all(d.day == 15 for d in paydays(frequency="monthly", monthly_day="15")))

    def test_no_payroll_produces_no_dates(self):
        self.assertEqual(paydays(frequency="none"), [])

    def test_window_starting_mid_year_spans_two_years(self):
        config = bb.validate_config(base(start_month=7, pay={"frequency": "monthly", "monthly_day": "1"}))
        days = bb.generate_paydays(config)
        self.assertEqual(days[0], date(2026, 7, 1))
        self.assertEqual(days[-1], date(2027, 6, 1))
        self.assertEqual(bb.period_label(config), "Jul 2026 - Jun 2027")


class SemimonthlyValidationTests(unittest.TestCase):
    def assert_invalid(self, days):
        with self.assertRaises(ValueError):
            paydays(frequency="semimonthly", semimonthly_days=days)

    def test_one_day_rejected(self):
        self.assert_invalid(["15"])

    def test_three_days_rejected(self):
        self.assert_invalid(["1", "15", "last"])

    def test_duplicate_days_rejected(self):
        self.assert_invalid(["15", "15"])

    def test_unknown_token_rejected(self):
        self.assert_invalid(["1", "20"])

    def test_missing_days_rejected(self):
        self.assert_invalid(None)


# ---------------------------------------------------------------- line expansion

class ExpansionTests(unittest.TestCase):
    def setUp(self):
        self.config = bb.validate_config(base())

    def expand(self, **item):
        line = bb.validate_line({"name": "x", **item}, "expenses", ("Bill", "Expense"))
        return bb.expand_line(line, self.config)

    def test_monthly(self):
        self.assertEqual(self.expand(amount=100, frequency="monthly"), [100.0] * 12)

    def test_weekly(self):
        self.assertAlmostEqual(self.expand(amount=12, frequency="weekly")[0], 52.0)

    def test_biweekly(self):
        self.assertAlmostEqual(self.expand(amount=12, frequency="biweekly")[0], 26.0)

    def test_semimonthly(self):
        self.assertEqual(self.expand(amount=50, frequency="semimonthly"), [100.0] * 12)

    def test_quarterly_with_month(self):
        values = self.expand(amount=300, frequency="quarterly", month=2)
        self.assertEqual([i for i, v in enumerate(values) if v], [1, 4, 7, 10])

    def test_yearly_with_month(self):
        values = self.expand(amount=600, frequency="yearly", month=12)
        self.assertEqual(values[11], 600)
        self.assertEqual(sum(values), 600)

    def test_one_time_defaults_to_first_month(self):
        values = self.expand(amount=75, frequency="one_time")
        self.assertEqual(values[0], 75)
        self.assertEqual(sum(values), 75)

    def test_month_mapping_follows_window_start(self):
        config = bb.validate_config(base(start_month=7))
        line = bb.validate_line({"name": "x", "amount": 10, "frequency": "yearly", "month": 1}, "income")
        self.assertEqual(bb.expand_line(line, config)[6], 10)  # Jan 2027 is window month 7

    def test_months_override(self):
        months = list(range(1, 13))
        self.assertEqual(self.expand(months=months, amount=999, frequency="monthly"), months)

    def test_months_wrong_length_rejected(self):
        with self.assertRaises(ValueError):
            self.expand(months=[1, 2, 3])


# ---------------------------------------------------------------- config validation

class ValidationTests(unittest.TestCase):
    def assert_invalid(self, raw):
        with self.assertRaises(ValueError):
            bb.validate_config(raw)

    def test_bad_mode(self):
        self.assert_invalid(base(mode="sepia"))

    def test_bad_accent(self):
        self.assert_invalid(base(accent="neon"))

    def test_bad_chart(self):
        self.assert_invalid(base(category_chart="radar"))

    def test_bad_budget_type(self):
        self.assert_invalid(base(budget_type="nonprofit"))

    def test_bad_pay_frequency(self):
        self.assert_invalid(base(pay={"frequency": "daily"}))

    def test_bad_line_frequency(self):
        self.assert_invalid(base(expenses=[{"name": "x", "amount": 1, "frequency": "hourly"}]))

    def test_bad_expense_type(self):
        self.assert_invalid(base(expenses=[{"name": "x", "amount": 1, "type": "Splurge"}]))

    def test_personal_type_rejected_in_business(self):
        self.assert_invalid(base(budget_type="business", expenses=[{"name": "x", "amount": 1, "type": "Bill"}]))

    def test_negative_amount(self):
        self.assert_invalid(base(expenses=[{"name": "x", "amount": -5}]))

    def test_negative_pay(self):
        self.assert_invalid(base(pay={"frequency": "monthly", "amount": -1}))

    def test_year_out_of_range(self):
        self.assert_invalid(base(year=1999))

    def test_missing_year(self):
        self.assert_invalid({"pay": {"frequency": "monthly"}})

    def test_bad_start_month(self):
        self.assert_invalid(base(start_month=13))

    def test_bad_date(self):
        self.assert_invalid(base(pay={"frequency": "weekly", "first_payday": "next friday"}))

    def test_bad_monthly_day(self):
        self.assert_invalid(base(pay={"frequency": "monthly", "monthly_day": "31"}))

    def test_missing_name(self):
        self.assert_invalid(base(income=[{"amount": 10}]))

    def test_defaults(self):
        config = bb.validate_config(base())
        self.assertEqual((config["budget_type"], config["mode"], config["category_chart"]),
                         ("personal", "light", "doughnut"))
        self.assertEqual(config["title"], "2026 Budget Planner")
        self.assertEqual(bb.validate_config(base(budget_type="business"))["title"], "2026 Business Budget")


# ---------------------------------------------------------------- summary math

class SummaryTests(unittest.TestCase):
    def test_personal_hand_computed(self):
        raw = base(pay={"frequency": "monthly", "amount": 3000, "monthly_day": "1"}, expenses=[
            {"name": "Rent", "category": "Housing", "type": "Bill", "amount": 1000},
            {"name": "Groceries", "category": "Food", "type": "Expense", "amount": 400}])
        summary = bb.compute_summary(bb.validate_config(raw))
        self.assertEqual(summary["annual_income"], 36000)
        self.assertEqual(summary["annual_fixed"], 12000)
        self.assertEqual(summary["annual_variable"], 4800)
        self.assertEqual(summary["annual_out"], 16800)
        self.assertEqual(summary["net"], 19200)
        self.assertAlmostEqual(summary["savings_rate"], 0.5333, places=4)
        self.assertEqual(len(summary["monthly"]), 12)

    def test_business_payroll_counts_as_expense(self):
        raw = base(budget_type="business", pay={"frequency": "semimonthly", "amount": 1000,
                                                "semimonthly_days": ["15", "last"]},
                   income=[{"name": "Sales", "amount": 10000}])
        summary = bb.compute_summary(bb.validate_config(raw))
        self.assertEqual(summary["annual_income"], 120000)
        self.assertEqual(summary["annual_fixed"], 24000)
        self.assertEqual(summary["net"], 96000)

    def test_zero_income_has_null_rate(self):
        self.assertIsNone(bb.compute_summary(bb.validate_config(base()))["savings_rate"])

    def test_three_paycheck_months_counted(self):
        raw = base(pay={"frequency": "biweekly", "amount": 100, "first_payday": "2026-01-02"})
        monthly = bb.compute_summary(bb.validate_config(raw))["monthly"]
        self.assertEqual(monthly[0]["in"], 300)  # Jan 2, 16, 30
        self.assertEqual(monthly[1]["in"], 200)


# ---------------------------------------------------------------- workbook structure

class WorkbookTests(TempDirCase):
    def example(self, name="example_config.json"):
        return json.loads((SCRIPTS / name).read_text(encoding="utf-8"))

    def test_sheet_order_with_trackers(self):
        _, wb = self.build(self.example())
        self.assertEqual(wb.sheetnames, ["Dashboard", "Income", "Paychecks", "Expenses",
                                         "Savings Goals", "Debt Tracker", "Lists"])

    def test_sheet_order_without_trackers(self):
        _, wb = self.build(base())
        self.assertEqual(wb.sheetnames, ["Dashboard", "Income", "Paychecks", "Expenses", "Lists"])

    def test_business_sheet_names(self):
        _, wb = self.build(self.example("example_business_config.json"))
        self.assertEqual(wb.sheetnames, ["Dashboard", "Revenue", "Payroll", "Expenses",
                                         "Reserve Goals", "Loans & Credit", "Lists"])

    def test_lists_hidden_and_dashboard_active(self):
        _, wb = self.build(base())
        self.assertEqual(wb["Lists"].sheet_state, "hidden")
        self.assertEqual(wb.active.title, "Dashboard")

    def test_dashboard_has_two_charts(self):
        _, wb = self.build(base())
        self.assertEqual(len(wb["Dashboard"]._charts), 2)

    def test_category_chart_matches_choice(self):
        expected = {"pie": ("PieChart", None), "doughnut": ("DoughnutChart", None),
                    "bar": ("BarChart", "bar"), "column": ("BarChart", "col")}
        for kind, (cls, direction) in expected.items():
            with self.subTest(kind=kind):
                bb_config = bb.validate_config(base(category_chart=kind))
                chart = bb.build_workbook(bb_config)["Dashboard"]._charts[1]
                self.assertEqual(type(chart).__name__, cls)
                if direction:
                    self.assertEqual(chart.type, direction)

    def test_dark_and_light_page_background(self):
        for mode in ("dark", "light"):
            with self.subTest(mode=mode):
                ws = bb.build_workbook(bb.validate_config(base(mode=mode)))["Dashboard"]
                self.assertEqual(ws["Z150"].fill.start_color.rgb[-6:], bb.MODES[mode]["bg"])

    def test_starter_rows_when_expenses_empty(self):
        _, wb = self.build(base())
        ws = wb["Expenses"]
        self.assertEqual(ws["B6"].value, "Rent / Mortgage")
        self.assertEqual(ws["D6"].value, "Bill")
        self.assertIsNone(ws["F6"].value)

    def test_business_starter_rows_and_payroll_row(self):
        _, wb = self.build(base(budget_type="business", pay={"frequency": "biweekly", "amount": 5000}))
        ws = wb["Expenses"]
        self.assertTrue(ws["B6"].value.startswith("Payroll"))
        self.assertEqual(ws["D6"].value, "Fixed")
        self.assertIn("SUMIFS('Payroll'!", ws["F6"].value)
        self.assertEqual(ws["B7"].value, "Office Rent / Lease")
        self.assertEqual(wb["Revenue"]["B6"].value, "Product Sales")

    def test_provided_expenses_written(self):
        raw = base(expenses=[{"name": "Gym", "category": "Health", "type": "Bill", "amount": 45, "due_day": 3}])
        _, wb = self.build(raw)
        ws = wb["Expenses"]
        self.assertEqual([ws["B6"].value, ws["C6"].value, ws["D6"].value, ws["E6"].value], ["Gym", "Health", "Bill", 3])
        self.assertEqual([ws.cell(row=6, column=c).value for c in range(6, 18)], [45] * 12)
        self.assertEqual(ws["R6"].value, '=IF(COUNT(F6:Q6)=0,"",SUM(F6:Q6))')

    def test_paycheck_rows_match_paydays(self):
        raw = base(pay={"frequency": "semimonthly", "amount": 1500, "semimonthly_days": ["1", "15"]})
        _, wb = self.build(raw)
        ws = wb["Paychecks"]
        dates = [ws.cell(row=r, column=3).value for r in range(6, 6 + 24)]
        self.assertTrue(all(dates))
        self.assertIsNone(ws.cell(row=6 + 24, column=3).value)  # blank spare rows follow
        self.assertEqual(ws["E6"].value, 1500)

    def test_income_paycheck_row_formula(self):
        _, wb = self.build(base())
        self.assertIn("SUMIFS('Paychecks'!", wb["Income"]["D6"].value)

    def test_dashboard_cells_are_formulas(self):
        _, wb = self.build(base())
        ws = wb["Dashboard"]
        for ref in ("B6", "D6", "F6", "H6", "K6", "M6", "C30", "D30", "E30", "H30", "L30"):
            self.assertTrue(str(ws[ref].value).startswith("="), ref)

    def test_kpis_reference_correct_total_columns(self):
        _, wb = self.build(base())
        ws = wb["Dashboard"]
        self.assertEqual((ws["B6"].value, ws["D6"].value, ws["F6"].value, ws["M6"].value),
                         ("=D42", "=G42", "=H42", "=C42"))
        self.assertEqual(ws["B42"].value, "TOTAL")

    def test_full_calc_on_load(self):
        _, wb = self.build(base())
        self.assertTrue(wb.calculation.fullCalcOnLoad)

    def test_dropdowns_on_category_and_type(self):
        _, wb = self.build(base())
        ranges = " ".join(str(dv.sqref) for dv in wb["Expenses"].data_validations.dataValidation)
        self.assertIn("C6:C", ranges)
        self.assertIn("D6:D", ranges)

    def test_custom_category_added_to_lists(self):
        _, wb = self.build(base(expenses=[{"name": "Vet", "category": "Pets", "amount": 50}]))
        values = [c.value for c in wb["Lists"]["A"]]
        self.assertIn("Pets", values)

    def test_chart_label_formats_unlinked(self):
        result, _ = self.build(base(category_chart="bar"))
        import zipfile
        with zipfile.ZipFile(result["file"]) as z:
            xml = z.read("xl/charts/chart2.xml").decode()
        self.assertNotRegex(xml, r'<numFmt formatCode="[^"]*"\s*/>')

    def test_print_areas_limited_to_tables(self):
        _, wb = self.build(base(features={"savings_goals": True, "debt_tracker": True}))
        for name in ("Dashboard", "Income", "Paychecks", "Expenses", "Savings Goals", "Debt Tracker"):
            with self.subTest(sheet=name):
                self.assertTrue(wb[name].print_area, name)
        self.assertIn("$G$", wb["Paychecks"].print_area)

    def test_default_output_names(self):
        self.assertEqual(bb.default_output_name(bb.validate_config(base())), "Budget Planner 2026.xlsx")
        self.assertEqual(bb.default_output_name(bb.validate_config(base(budget_type="business", start_month=7))),
                         "Business Budget Jul 2026-Jun 2027.xlsx")


# ---------------------------------------------------------------- CLI

class CliTests(TempDirCase):
    def run_cli(self, config_path: Path):
        out = self.tmp / "cli.xlsx"
        return subprocess.run([sys.executable, str(SCRIPT), "--config", str(config_path), "--out", str(out)],
                              capture_output=True, text=True)

    def test_valid_config_exit_zero_and_json(self):
        proc = self.run_cli(SCRIPTS / "example_config.json")
        self.assertEqual(proc.returncode, 0, proc.stderr)
        result = json.loads(proc.stdout)
        self.assertTrue(Path(result["file"]).exists())
        self.assertEqual(result["paychecks"], 26)

    def test_invalid_config_exit_two(self):
        bad = self.tmp / "bad.json"
        bad.write_text(json.dumps(base(mode="sepia")), encoding="utf-8")
        proc = self.run_cli(bad)
        self.assertEqual(proc.returncode, 2)
        self.assertIn("ERROR:", proc.stderr)

    def test_missing_config_exit_two(self):
        proc = self.run_cli(self.tmp / "nope.json")
        self.assertEqual(proc.returncode, 2)
        self.assertIn("not found", proc.stderr)

    def test_config_with_utf8_bom_accepted(self):
        bom = self.tmp / "bom.json"
        bom.write_bytes(b"\xef\xbb\xbf" + json.dumps(base()).encode("utf-8"))
        self.assertEqual(self.run_cli(bom).returncode, 0)

    def test_malformed_json_exit_two(self):
        bad = self.tmp / "broken.json"
        bad.write_text("{year: 2026", encoding="utf-8")
        self.assertEqual(self.run_cli(bad).returncode, 2)


if __name__ == "__main__":
    unittest.main()
