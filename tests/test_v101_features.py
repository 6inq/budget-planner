"""Tests for 1.0.1: credit line, saved settings, workbook round trip, CSV statement import, privacy.

Run from the repository root:
    python -m unittest discover -s tests -v
"""

import ast
import importlib.util
import json
import subprocess
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

SCRIPTS = Path(__file__).resolve().parents[1] / "skills" / "budget-planner" / "scripts"
sys.path.insert(0, str(SCRIPTS))


def load(name: str):
    spec = importlib.util.spec_from_file_location(name, SCRIPTS / f"{name}.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


bb = load("build_budget")
rw = load("read_workbook")
it = load("import_transactions")


def base(**overrides) -> dict:
    raw = {"year": 2026, "pay": {"frequency": "monthly", "amount": 0}}
    raw.update(overrides)
    return raw


def example(name="example_config.json") -> dict:
    return json.loads((SCRIPTS / name).read_text(encoding="utf-8"))


class TempDirCase(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def build(self, raw: dict, name="out.xlsx", **kwargs) -> Path:
        return Path(bb.build_from_config(raw, str(self.tmp / name), **kwargs)["file"])

    def write_csv(self, text: str, name="statement.csv") -> Path:
        path = self.tmp / name
        path.write_text(text, encoding="utf-8")
        return path


# ---------------------------------------------------------------- credit line

class CreditTests(TempDirCase):
    def test_every_visible_sheet_has_credit_and_link(self):
        wb = load_workbook(self.build(base(features={"savings_goals": True, "debt_tracker": True})))
        for ws in wb.worksheets:
            if ws.sheet_state == "hidden":
                continue
            with self.subTest(sheet=ws.title):
                self.assertEqual(ws["B4"].value, bb.CREDIT_TEXT)
                self.assertEqual(ws["B4"].hyperlink.target, bb.CREDIT_URL)

    def test_credit_is_unobtrusive(self):
        ws = load_workbook(self.build(base(mode="dark")))["Dashboard"]
        font = ws["B4"].font
        self.assertLessEqual(font.sz, 8)
        self.assertEqual(font.color.rgb[-6:], bb.MODES["dark"]["muted"])

    def test_business_workbook_has_credit(self):
        wb = load_workbook(self.build(example("example_business_config.json")))
        self.assertEqual(wb["Revenue"]["B4"].value, bb.CREDIT_TEXT)

    def test_print_footer_and_file_properties(self):
        wb = load_workbook(self.build(base()))
        self.assertEqual(wb["Expenses"].oddFooter.right.text, bb.CREDIT_TEXT)
        self.assertEqual(wb.properties.creator, bb.CREDIT_TEXT)


# ---------------------------------------------------------------- settings + explicit paychecks

class SettingsTests(TempDirCase):
    def test_settings_saved_without_financial_data(self):
        wb = load_workbook(self.build(example()))
        settings = json.loads(wb["Lists"][bb.SETTINGS_CELL].value)
        self.assertEqual((settings["mode"], settings["category_chart"]), ("dark", "doughnut"))
        self.assertEqual(settings["pay"]["frequency"], "biweekly")
        for key in ("income", "expenses", "savings_goals", "debts"):
            self.assertNotIn(key, settings)
        self.assertNotIn("paychecks", settings["pay"])

    def test_explicit_paychecks_replace_schedule(self):
        raw = base(pay={"frequency": "biweekly", "amount": 1000, "paychecks": [
            {"date": "2026-01-09", "amount": 1000}, {"date": "2026-01-23", "amount": 1500, "note": "OT"}]})
        config = bb.validate_config(raw)
        self.assertEqual(bb.compute_summary(config)["annual_income"], 2500)
        ws = load_workbook(self.build(raw))["Paychecks"]
        self.assertEqual((ws["E7"].value, ws["F7"].value), (1500, "OT"))

    def test_paycheck_outside_window_rejected(self):
        raw = base(pay={"frequency": "monthly", "paychecks": [{"date": "2025-12-31", "amount": 5}]})
        with self.assertRaises(ValueError):
            bb.validate_config(raw)

    def test_backup_created_only_when_file_exists(self):
        first = bb.build_from_config(base(), str(self.tmp / "b.xlsx"), backup=True)
        self.assertNotIn("backup", first)
        second = bb.build_from_config(base(), str(self.tmp / "b.xlsx"), backup=True)
        self.assertTrue(Path(second["backup"]).exists())
        self.assertTrue(second["backup"].endswith("b (backup).xlsx"))


# ---------------------------------------------------------------- workbook round trip

class RoundTripTests(TempDirCase):
    def rebuild(self, path: Path) -> dict:
        config = rw.read_workbook(str(path))
        return bb.build_from_config(config, str(self.tmp / "rebuilt.xlsx"))

    def test_round_trip_keeps_every_total(self):
        for name in ("example_config.json", "example_business_config.json"):
            with self.subTest(config=name):
                original = bb.build_from_config(example(name), str(self.tmp / f"orig-{name}.xlsx"))
                rebuilt = self.rebuild(Path(original["file"]))
                for key in ("annual_income", "annual_out", "net", "paychecks"):
                    self.assertEqual(original[key], rebuilt[key], key)
                self.assertEqual(original["monthly"], rebuilt["monthly"])

    def test_round_trip_keeps_settings(self):
        config = rw.read_workbook(str(self.build(example("example_business_config.json"))))
        self.assertEqual((config["budget_type"], config["mode"], config["accent"]), ("business", "light", "teal"))
        self.assertEqual(config["pay"]["semimonthly_days"], ["15", "last"])

    def test_user_edits_in_excel_are_kept(self):
        path = self.build(base(pay={"frequency": "monthly", "amount": 3000}))
        wb = load_workbook(path)
        wb["Expenses"]["F6"] = 1200            # Rent / Mortgage, January
        wb["Expenses"]["B30"] = "Dog Walker"   # a brand new row typed by the user
        wb["Expenses"]["C30"] = "Pets"
        wb["Expenses"]["D30"] = "expense"      # lowercase type still recognized
        wb["Expenses"]["G30"] = 80
        wb["Paychecks"]["E7"] = 3500           # February paycheck with a bonus
        wb["Paychecks"]["F7"] = "Bonus"
        wb["Income"]["B7"] = "Etsy"
        wb["Income"]["D7"] = "250"             # number typed as text
        wb.save(path)
        config = rw.read_workbook(str(path))
        rent = next(e for e in config["expenses"] if e["name"] == "Rent / Mortgage")
        walker = next(e for e in config["expenses"] if e["name"] == "Dog Walker")
        self.assertEqual(rent["months"][0], 1200)
        self.assertEqual((walker["category"], walker["type"], walker["months"][1]), ("Pets", "Expense", 80))
        self.assertEqual((config["pay"]["paychecks"][1]["amount"], config["pay"]["paychecks"][1]["note"]),
                         (3500, "Bonus"))
        self.assertEqual(config["income"][0]["months"][0], 250)
        summary = bb.compute_summary(bb.validate_config(config))
        self.assertEqual(summary["annual_income"], 3000 * 12 + 500 + 250)

    def test_starter_rows_survive_round_trip(self):
        config = rw.read_workbook(str(self.build(base())))
        self.assertIn("Groceries", [e["name"] for e in config["expenses"]])

    def test_business_payroll_row_not_duplicated(self):
        config = rw.read_workbook(str(self.build(example("example_business_config.json"))))
        self.assertFalse(any(e["name"].startswith("Payroll") for e in config["expenses"]))

    def test_trackers_read_back(self):
        config = rw.read_workbook(str(self.build(example())))
        self.assertEqual(config["savings_goals"][0]["target_date"], "2026-12-31")
        self.assertEqual(config["debts"][0]["apr"], 22.9)

    def test_text_in_number_cell_is_reported(self):
        path = self.build(base())
        wb = load_workbook(path)
        wb["Expenses"]["F7"] = "about forty"
        wb.save(path)
        with self.assertRaisesRegex(ValueError, "Expenses!F7"):
            rw.read_workbook(str(path))

    def test_formula_without_saved_result_is_reported(self):
        path = self.build(base())
        wb = load_workbook(path)
        wb["Expenses"]["F7"] = "=20*2"
        wb.save(path)
        with self.assertRaisesRegex(ValueError, "Open the workbook in Excel"):
            rw.read_workbook(str(path))

    def test_foreign_workbook_rejected(self):
        from openpyxl import Workbook
        path = self.tmp / "other.xlsx"
        Workbook().save(path)
        with self.assertRaisesRegex(ValueError, "no saved Budget Planner settings"):
            rw.read_workbook(str(path))

    def test_cli_writes_config(self):
        path = self.build(example())
        out = self.tmp / "cfg.json"
        proc = subprocess.run([sys.executable, str(SCRIPTS / "read_workbook.py"), "--workbook", str(path),
                               "--out", str(out)], capture_output=True, text=True)
        self.assertEqual(proc.returncode, 0, proc.stderr)
        self.assertEqual(json.loads(proc.stdout)["paychecks"], 26)
        self.assertTrue(json.loads(out.read_text(encoding="utf-8"))["expenses"])

    def test_cli_missing_workbook(self):
        proc = subprocess.run([sys.executable, str(SCRIPTS / "read_workbook.py"), "--workbook",
                               str(self.tmp / "nope.xlsx")], capture_output=True, text=True)
        self.assertEqual(proc.returncode, 2)


# ---------------------------------------------------------------- CSV statement import

CHECKING_CSV = """Account Number: XXXX1234
Statement Period: 01/01/2026 - 03/31/2026

Date,Description,Amount,Balance
01/01/2026,ACH DEBIT RENT PAYMENT GREENVIEW APTS,-1650.00,3000.00
01/09/2026,ACME CORP PAYROLL,2150.00,5150.00
01/12/2026,POS NETFLIX.COM 866-579-7172,-15.49,5134.51
01/15/2026,KROGER #1234 COLUMBUS OH,-84.12,5050.39
01/28/2026,KROGER #5678 COLUMBUS OH,-61.30,4989.09
02/01/2026,ACH DEBIT RENT PAYMENT GREENVIEW APTS,-1650.00,3339.09
02/12/2026,POS NETFLIX.COM 866-579-7172,-15.49,3323.60
02/20/2026,KROGER #1234 COLUMBUS OH,"-$1,102.40",2221.20
03/01/2026,ACH DEBIT RENT PAYMENT GREENVIEW APTS,(1650.00),571.20
03/12/2026,POS NETFLIX.COM 866-579-7172,-15.49,555.71
not a date,junk row,12,0
"""


class StatementImportTests(TempDirCase):
    def run_import(self, text: str, **kwargs) -> dict:
        return it.run([str(self.write_csv(text))], **kwargs)

    def merchant(self, result: dict, direction: str, name: str) -> dict:
        return next(g for g in result[direction] if g["merchant"] == name)

    def test_preamble_skipped_and_header_found(self):
        result = self.run_import(CHECKING_CSV)
        self.assertEqual(result["transactions"], 10)
        self.assertEqual(result["skipped_rows"], 1)
        self.assertEqual(result["months_covered"], ["2026-01", "2026-02", "2026-03"])

    def test_merchant_normalization_groups_store_numbers(self):
        result = self.run_import(CHECKING_CSV)
        kroger = self.merchant(result, "money_out", "Kroger Columbus Oh")
        self.assertEqual(kroger["count"], 3)
        self.assertEqual(kroger["total"], 1247.82)

    def test_exact_totals_with_currency_and_parentheses(self):
        result = self.run_import(CHECKING_CSV)
        rent = self.merchant(result, "money_out", "Rent Payment Greenview Apts")
        self.assertEqual(rent["total"], 4950.0)
        self.assertEqual(rent["months"], {"2026-01": 1650.0, "2026-02": 1650.0, "2026-03": 1650.0})

    def test_recurring_detection(self):
        result = self.run_import(CHECKING_CSV)
        self.assertTrue(self.merchant(result, "money_out", "Rent Payment Greenview Apts")["recurring"])
        self.assertTrue(self.merchant(result, "money_out", "Netflix Com")["recurring"])
        self.assertFalse(self.merchant(result, "money_out", "Kroger Columbus Oh")["recurring"])
        self.assertEqual(self.merchant(result, "money_out", "Rent Payment Greenview Apts")["typical_day"], 1)

    def test_money_in_separated(self):
        result = self.run_import(CHECKING_CSV)
        self.assertEqual(result["money_in"][0]["merchant"], "Acme Corp Payroll")
        self.assertEqual(result["money_in_total"], 2150.0)

    def test_window_arrays(self):
        result = self.run_import(CHECKING_CSV, start="2026-01")
        rent = self.merchant(result, "money_out", "Rent Payment Greenview Apts")
        self.assertEqual(rent["window_months"], [1650.0, 1650.0, 1650.0] + [0.0] * 9)

    def test_debit_credit_columns(self):
        text = "Posted Date,Payee,Debit,Credit\n2026-01-05,Electric Co,120.50,\n2026-01-06,Refund,,20.00\n"
        result = self.run_import(text)
        self.assertEqual(result["money_out_total"], 120.5)
        self.assertEqual(result["money_in_total"], 20.0)

    def test_type_column_sets_direction(self):
        text = "Date,Description,Amount,Type\n01/03/2026,Gym,45.00,Debit\n01/04/2026,Transfer In,500,Credit\n"
        result = self.run_import(text)
        self.assertEqual(result["money_out"][0]["merchant"], "Gym")
        self.assertEqual(result["money_in_total"], 500)

    def test_invert_for_card_exports(self):
        text = "Transaction Date,Description,Amount\n01/10/2026,AMAZON MKTPLACE,54.99\n01/20/2026,PAYMENT THANK YOU,-54.99\n"
        result = self.run_import(text, invert=True)
        self.assertEqual(result["money_out"][0]["merchant"], "Amazon Mktplace")

    def test_day_first_detected(self):
        text = "Date,Description,Amount\n25/01/2026,Tesco,-30.00\n03/02/2026,Tesco,-20.00\n"
        result = self.run_import(text)
        self.assertEqual(result["months_covered"], ["2026-01", "2026-02"])

    def test_semicolon_delimiter_and_bom(self):
        path = self.tmp / "eu.csv"
        path.write_bytes(b"\xef\xbb\xbf" + "Date;Description;Amount\n2026-01-02;Rent;-900\n".encode("utf-8"))
        self.assertEqual(it.run([str(path)])["money_out_total"], 900)

    def test_multiple_files_combined(self):
        a = self.write_csv("Date,Description,Amount\n01/02/2026,Gym,-45\n", "a.csv")
        b = self.write_csv("Date,Description,Amount\n02/02/2026,Gym,-45\n", "b.csv")
        result = it.run([str(a), str(b)])
        self.assertEqual(result["money_out"][0]["count"], 2)
        self.assertEqual(len(result["sources"]), 2)

    def test_masked_card_numbers_removed(self):
        self.assertEqual(it.normalize_merchant("CHECKCARD 0112 SHELL OIL 57444 XXXX1234"), "Shell Oil")

    def test_missing_columns_rejected(self):
        with self.assertRaisesRegex(ValueError, "header row"):
            self.run_import("Foo,Bar\n1,2\n")

    def test_empty_file_rejected(self):
        with self.assertRaises(ValueError):
            self.run_import("")

    def test_bad_start_rejected(self):
        with self.assertRaises(ValueError):
            self.run_import(CHECKING_CSV, start="January")

    def test_cli(self):
        path = self.write_csv(CHECKING_CSV)
        proc = subprocess.run([sys.executable, str(SCRIPTS / "import_transactions.py"), str(path), "--start", "2026-01"],
                              capture_output=True, text=True)
        self.assertEqual(proc.returncode, 0, proc.stderr)
        self.assertIn("window", json.loads(proc.stdout))
        bad = subprocess.run([sys.executable, str(SCRIPTS / "import_transactions.py"), str(self.tmp / "none.csv")],
                             capture_output=True, text=True)
        self.assertEqual(bad.returncode, 2)

    def test_output_never_contains_account_number_lines(self):
        result = json.dumps(self.run_import(CHECKING_CSV))
        self.assertNotIn("XXXX1234", result)
        self.assertNotIn("Account Number", result)


# ---------------------------------------------------------------- privacy guarantee

class PrivacyTests(unittest.TestCase):
    NETWORK_MODULES = {"socket", "urllib", "http", "requests", "httpx", "ftplib", "smtplib", "aiohttp",
                       "urllib3", "websocket", "telnetlib", "xmlrpc"}

    def test_scripts_import_no_network_modules(self):
        for script in SCRIPTS.glob("*.py"):
            tree = ast.parse(script.read_text(encoding="utf-8"))
            imported = set()
            for node in ast.walk(tree):
                if isinstance(node, ast.Import):
                    imported |= {alias.name.split(".")[0] for alias in node.names}
                elif isinstance(node, ast.ImportFrom) and node.module:
                    imported.add(node.module.split(".")[0])
            with self.subTest(script=script.name):
                self.assertFalse(imported & self.NETWORK_MODULES, imported & self.NETWORK_MODULES)


if __name__ == "__main__":
    unittest.main()
