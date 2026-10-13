"""Read a Budget Planner workbook back into a config file.

Usage:
    python read_workbook.py --workbook "Budget Planner 2026.xlsx" --out budget-config.json

Updating a budget means: read the workbook (so edits the user typed directly in
Excel are kept), merge in the new numbers, then rebuild with build_budget.py.
Everything happens locally; this script never opens a network connection.
"""

import argparse
import json
import sys
import warnings
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import build_budget as bb  # noqa: E402
from profiles import PROFILES  # noqa: E402

DATE_FORMATS = ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%d/%m/%Y")

# Excel-saved files carry extensions openpyxl cannot parse; harmless here because we only read values.
warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")


class SheetReader:
    """Reads input cells, falling back to Excel's saved result when a cell holds a formula."""

    def __init__(self, formulas_ws, values_ws):
        self.f, self.v, self.title = formulas_ws, values_ws, formulas_ws.title

    def ref(self, row: int, col: int) -> str:
        return f"{self.title}!{get_column_letter(col)}{row}"

    def value(self, row: int, col: int):
        raw = self.f.cell(row=row, column=col).value
        if isinstance(raw, str) and raw.startswith("="):
            cached = self.v.cell(row=row, column=col).value
            if cached is None:
                raise ValueError(f"{self.ref(row, col)} has a formula with no saved result. "
                                 "Open the workbook in Excel, save it, and try again.")
            return cached
        return raw

    def text(self, row: int, col: int) -> str:
        value = self.value(row, col)
        return "" if value is None else str(value).strip()

    def number(self, row: int, col: int) -> float:
        value = self.value(row, col)
        if value is None or value == "":
            return 0.0
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            return float(value)
        cleaned = str(value).replace(",", "").strip().strip("$€£ ")
        try:
            return float(cleaned)
        except ValueError:
            raise ValueError(f"{self.ref(row, col)} should be a number but contains {value!r}") from None

    def day(self, row: int, col: int):
        value = self.value(row, col)
        if value in (None, ""):
            return None
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        for fmt in DATE_FORMATS:
            try:
                return datetime.strptime(str(value).strip(), fmt).date()
            except ValueError:
                continue
        raise ValueError(f"{self.ref(row, col)} should be a date but contains {value!r}")

    def total_row(self, col: int) -> int:
        """Our TOTAL rows are the only cells holding =SUM(<col>6:...) in that column."""
        prefix = f"=SUM({get_column_letter(col)}{bb.FIRST_ROW}:"
        for row in range(bb.FIRST_ROW, self.f.max_row + 1):
            value = self.f.cell(row=row, column=col).value
            if isinstance(value, str) and value.upper().startswith(prefix):
                return row
        raise ValueError(f"could not find the TOTAL row on the '{self.title}' sheet. Was it deleted?")


def open_reader(wb_formulas, wb_values, name: str) -> SheetReader:
    return SheetReader(wb_formulas[name], wb_values[name])


def read_settings(wb) -> dict:
    if "Lists" not in wb.sheetnames or not wb["Lists"][bb.SETTINGS_CELL].value:
        raise ValueError("this workbook has no saved Budget Planner settings. It was made by another tool "
                         "or an older version; answer the questions again to build a new one.")
    settings = json.loads(wb["Lists"][bb.SETTINGS_CELL].value)
    settings.pop("format", None)
    return settings


def read_month_lines(sheet: SheetReader, start: int, month_col: int, expense_types=None) -> list[dict]:
    items = []
    for row in range(start, sheet.total_row(month_col)):
        name = sheet.text(row, 2)
        months = [sheet.number(row, month_col + i) for i in range(12)]
        if not name and not any(months):
            continue
        item = {"name": name or f"Row {row}", "category": sheet.text(row, 3) or "Other", "months": months}
        if expense_types:
            item["type"] = normalize_type(sheet.text(row, 4), expense_types)
            due = sheet.number(row, 5)
            item["due_day"] = int(due) if 1 <= due <= 31 else None
            item["paid"] = [sheet.text(row, bb.PAID_COL + i).lower() == bb.PAID_MARK.lower() for i in range(12)]
        items.append(item)
    return items


def normalize_type(value: str, types: tuple) -> str:
    matches = [t for t in types if t.lower() == value.lower()]
    return matches[0] if matches else types[1]


def read_paychecks(sheet: SheetReader) -> list[dict]:
    entries = []
    for row in range(bb.FIRST_ROW, sheet.total_row(5)):
        day = sheet.day(row, 3)
        if day:
            entries.append({"date": day.isoformat(), "amount": sheet.number(row, 5), "note": sheet.text(row, 6)})
    return entries


def read_goals(sheet: SheetReader) -> list[dict]:
    goals = []
    for row in range(bb.FIRST_ROW, sheet.total_row(3)):
        name, target, saved = sheet.text(row, 2), sheet.number(row, 3), sheet.number(row, 4)
        if name or target or saved:
            target_date = sheet.day(row, 7)
            goals.append({"name": name or "Goal", "target": target, "saved": saved,
                          "target_date": target_date.isoformat() if target_date else None})
    return goals


def read_debts(sheet: SheetReader) -> list[dict]:
    debts = []
    for row in range(bb.FIRST_ROW, sheet.total_row(3)):
        name, balance = sheet.text(row, 2), sheet.number(row, 3)
        if name or balance:
            debts.append({"name": name or "Debt", "balance": balance,
                          "apr": sheet.number(row, 4), "min_payment": sheet.number(row, 5)})
    return debts


def read_payoff_plan(sheet: SheetReader, loan_type: str) -> dict:
    plan = {"loan_type": loan_type, "loan_name": sheet.text(bb.plan_row("loan_name"), 3)}
    for key in bb.PLAN_AMOUNT_KEYS:
        raw = sheet.value(bb.plan_row(key), 3)
        plan[key] = None if raw in (None, "") else sheet.number(bb.plan_row(key), 3)
    start = sheet.day(bb.plan_row("start"), 3)
    plan["start"] = start.isoformat() if start else None
    first = bb.PLAN_HEADER_ROW + 1
    extra = [None if sheet.value(first + n, 7) in (None, "") else sheet.number(first + n, 7)
             for n in range(bb.PLAN_HORIZONS[loan_type])]
    while extra and extra[-1] is None:
        extra.pop()
    plan["extra_paid"] = extra
    return plan


def read_workbook(path: str) -> dict:
    wb_formulas = load_workbook(path)
    wb_values = load_workbook(path, data_only=True)
    config = read_settings(wb_formulas)
    loan_type = config.pop("payoff_loan_type", "car")
    profile = PROFILES[config.get("budget_type", "personal")]
    sheets = profile["sheets"]
    pay_is_income = profile["pay_is_income"]
    income = open_reader(wb_formulas, wb_values, sheets["income"])
    expenses = open_reader(wb_formulas, wb_values, sheets["expenses"])
    config["income"] = read_month_lines(income, bb.FIRST_ROW + (1 if pay_is_income else 0), bb.INCOME_MONTH_COL)
    config["expenses"] = read_month_lines(expenses, bb.FIRST_ROW + (0 if pay_is_income else 1),
                                          bb.EXPENSE_MONTH_COL, profile["types"])
    config["pay"]["paychecks"] = read_paychecks(open_reader(wb_formulas, wb_values, sheets["pay"]))
    has_goals, has_debts = sheets["goals"] in wb_formulas.sheetnames, sheets["debts"] in wb_formulas.sheetnames
    has_plan = sheets["payoff"] in wb_formulas.sheetnames
    config["features"] = {"savings_goals": has_goals, "debt_tracker": has_debts, "payoff_plan": has_plan}
    if has_plan:
        config["payoff_plan"] = read_payoff_plan(open_reader(wb_formulas, wb_values, sheets["payoff"]), loan_type)
    config["savings_goals"] = read_goals(open_reader(wb_formulas, wb_values, sheets["goals"])) if has_goals else []
    config["debts"] = read_debts(open_reader(wb_formulas, wb_values, sheets["debts"])) if has_debts else []
    bb.validate_config(config)  # fail here, with a clear message, rather than at rebuild time
    return config


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="Read a Budget Planner workbook back into a config file.")
    parser.add_argument("--workbook", required=True, help="Path to the existing .xlsx")
    parser.add_argument("--out", default="budget-config.json", help="Config file to write")
    args = parser.parse_args(argv)
    try:
        config = read_workbook(args.workbook)
    except FileNotFoundError:
        print(f"ERROR: workbook not found: {args.workbook}", file=sys.stderr)
        return 2
    except (ValueError, KeyError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2
    Path(args.out).write_text(json.dumps(config, indent=2), encoding="utf-8")
    print(json.dumps({"config": str(Path(args.out).resolve()), "budget_type": config["budget_type"],
                      "income_rows": len(config["income"]), "expense_rows": len(config["expenses"]),
                      "paychecks": len(config["pay"]["paychecks"]), "savings_goals": len(config["savings_goals"]),
                      "debts": len(config["debts"])}, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
